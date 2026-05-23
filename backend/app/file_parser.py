"""
文件内容提取模块
支持: .md, .txt, .docx, .pdf
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".md", ".txt", ".docx", ".pdf"}

MAX_CHARS = 30000  # 提取文本最大字符数，防止超大文件


def extract_text(file_path: Path, filename: str) -> str:
    """根据文件扩展名提取纯文本内容"""
    ext = Path(filename).suffix.lower()

    if ext in (".md", ".txt"):
        return _extract_plain(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    elif ext == ".pdf":
        return _extract_pdf(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}，支持: {', '.join(SUPPORTED_EXTENSIONS)}")


def _extract_plain(file_path: Path) -> str:
    """md/txt 直接读文本"""
    text = file_path.read_text(encoding="utf-8", errors="replace")
    return text[:MAX_CHARS]


def _extract_docx(file_path: Path) -> str:
    """从 docx 提取段落文本"""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("需要安装 python-docx: pip install python-docx")

    doc = Document(str(file_path))
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # 保留标题层级信息
            style_name = para.style.name if para.style else ""
            if "Heading 1" in style_name:
                paragraphs.append(f"# {text}")
            elif "Heading 2" in style_name:
                paragraphs.append(f"## {text}")
            elif "Heading 3" in style_name:
                paragraphs.append(f"### {text}")
            else:
                paragraphs.append(text)

    # 也提取表格内容
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cells)
            if row_text.strip("| "):
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs)
    return full_text[:MAX_CHARS]


def _extract_pdf(file_path: Path) -> str:
    """从 PDF 提取文本（优先 pymupdf，备选 pdfplumber）"""
    text_parts = []

    # 尝试 pymupdf (fitz)
    try:
        import fitz  # pymupdf
        doc = fitz.open(str(file_path))
        for page in doc:
            page_text = page.get_text()
            if page_text.strip():
                text_parts.append(page_text.strip())
        doc.close()
        full_text = "\n\n".join(text_parts)
        if full_text.strip():
            return full_text[:MAX_CHARS]
    except ImportError:
        pass
    except Exception as e:
        logger.warning("pymupdf 提取失败: %s", e)

    # 备选 pdfplumber
    try:
        import pdfplumber
        with pdfplumber.open(str(file_path)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    text_parts.append(page_text.strip())
        full_text = "\n\n".join(text_parts)
        return full_text[:MAX_CHARS]
    except ImportError:
        raise ImportError("需要安装 pymupdf 或 pdfplumber 来处理 PDF 文件")
    except Exception as e:
        raise RuntimeError(f"PDF 文本提取失败: {e}")
